"""Adapter: write derived site codes into a copy of the Master Planner.

Why this exists
---------------
Roughly a fifth of the Master Planner's customer columns carry no account number
in their header. The converter still gives those sites a code, but it does so by
hashing the header text — which means the code only survives as long as the header
text does. Correct a typo in a header and the hash changes, the code changes, and
the site silently stops matching between the manual plan and the optimizer input.

Writing the codes into the headers removes that fragility: the code becomes a
literal part of the workbook, so it survives edits, reordering and renaming.

This produces a **copy**. The planner's own workbook is never modified in place.
Formulas are preserved (the workbook is opened without ``data_only``), so the
returned file is a working Master Planner and not a values-only snapshot.

Only this layer touches openpyxl.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from typing import Sequence

import openpyxl

from domain.site_derivation import DerivedSite
from io_adapters.master_planner_parser import HEADER_ROW

# Separator between the code and the existing header text.
CODE_SEPARATOR = "  "


@dataclass
class AnnotationResult:
    """Outcome of writing codes into a Master Planner copy."""

    xlsx_bytes: bytes | None = None
    written: list[tuple[str, str]] = field(default_factory=list)   # (code, column)
    already_present: list[tuple[str, str]] = field(default_factory=list)
    skipped: list[tuple[str, str]] = field(default_factory=list)   # (code, reason)

    @property
    def written_count(self) -> int:
        return len(self.written)


class MasterPlannerAnnotator:
    """Writes generated site codes into Master Planner column headers.

    Implements :class:`services.ports.MasterPlannerAnnotatorPort`.
    """

    def annotate(
        self,
        source: bytes,
        sites: Sequence[DerivedSite],
        sheet: str = "Schedule",
    ) -> AnnotationResult:
        """Return a copy of the workbook with generated codes in the headers.

        Parameters
        ----------
        source : bytes
            The original Master Planner workbook.
        sites : Sequence[DerivedSite]
            Derived sites. Only those whose code was generated (no account
            number in the header) are written; sites that already carry an
            account number are left untouched, because their header is already
            the shared key.
        sheet : str
            Worksheet holding the schedule grid.

        Raises
        ------
        ValueError
            If the sheet is missing or has no header row.
        """
        result = AnnotationResult()

        # data_only must stay False: opening with cached values and saving would
        # strip every formula out of the planner's workbook.
        wb = openpyxl.load_workbook(io.BytesIO(source))
        try:
            if sheet not in wb.sheetnames:
                raise ValueError(
                    f"Sheet '{sheet}' not found. Available: {wb.sheetnames}"
                )
            ws = wb[sheet]
            if ws.max_row < HEADER_ROW:
                raise ValueError("Master Planner sheet has no header row.")

            for site in sites:
                if site.code_source != "generated":
                    continue
                if not site.column_letter:
                    result.skipped.append(
                        (site.site_code, "no column reference was recorded")
                    )
                    continue

                cell = ws[f"{site.column_letter}{HEADER_ROW}"]
                current = "" if cell.value is None else str(cell.value)

                if current.strip().startswith(site.site_code):
                    result.already_present.append(
                        (site.site_code, site.column_letter)
                    )
                    continue

                cell.value = f"{site.site_code}{CODE_SEPARATOR}{current.strip()}"
                result.written.append((site.site_code, site.column_letter))

            buffer = io.BytesIO()
            wb.save(buffer)
            result.xlsx_bytes = buffer.getvalue()
        finally:
            wb.close()

        return result


def annotate_master_planner(
    source: bytes,
    sites: Sequence[DerivedSite],
    sheet: str = "Schedule",
) -> AnnotationResult:
    """Module-level convenience wrapper around :class:`MasterPlannerAnnotator`."""
    return MasterPlannerAnnotator().annotate(source, sites, sheet)
