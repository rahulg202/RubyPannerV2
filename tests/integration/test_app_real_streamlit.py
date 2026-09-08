"""Render the app against the real Streamlit API, not the test stub.

Why this exists
---------------
``tests/conftest.py`` replaces Streamlit with a ``MagicMock`` so the UI can be
imported without a runtime. That stub accepts *any* argument, which is exactly
how ``st.data_editor(..., width="stretch")`` passed 574 tests and then crashed in
production with ``TypeError: 'str' object cannot be interpreted as an integer``:
the deployed Streamlit accepted the literal, the locally-installed one wanted an
int, and the stub cared about neither.

These tests run the app through ``streamlit.testing.v1.AppTest``, which executes
the script against the genuine widget API and reports any exception the script
raised. They run in a **subprocess** so the conftest stub — installed globally at
import time — cannot leak in.

Scope: rendering only. No uploads, no solver runs. That is enough to catch widget
signatures drifting away from the pinned Streamlit version, which is the failure
this guards against.
"""

from __future__ import annotations

import subprocess
import sys
from pathlib import Path

import pytest

REPO_ROOT = Path(__file__).resolve().parents[2]

# Marker the child prints on success; anything else means the render failed.
OK = "APPTEST_OK"


def _run(script: str) -> subprocess.CompletedProcess:
    """Execute a snippet in a clean interpreter rooted at the repo."""
    return subprocess.run(
        [sys.executable, "-c", script],
        cwd=REPO_ROOT,
        capture_output=True,
        text=True,
        timeout=600,
    )


def _assert_rendered(result: subprocess.CompletedProcess) -> None:
    combined = f"{result.stdout}\n{result.stderr}"
    assert result.returncode == 0 and OK in result.stdout, combined


# Shared preamble: build the AppTest and a reporter for script exceptions.
_HEADER = f"""
import sys
from streamlit.testing.v1 import AppTest

def report(at):
    if at.exception:
        for exc in at.exception:
            sys.stderr.write(str(getattr(exc, "value", exc)) + "\\n")
            sys.stderr.write(str(getattr(exc, "stack_trace", "")) + "\\n")
        raise SystemExit(1)
    print("{OK}")
"""


def test_app_renders_every_tab():
    """A cold start must render all five tabs without raising."""
    _assert_rendered(_run(_HEADER + """
at = AppTest.from_file("app.py", default_timeout=300)
at.run()
report(at)
"""))


def test_app_renders_with_the_settings_expanders_touched():
    """Non-default settings must not break any widget."""
    _assert_rendered(_run(_HEADER + """
from datetime import date

at = AppTest.from_file("app.py", default_timeout=300)
at.session_state["cfg_use_reference"] = True
at.session_state["cfg_reference_week_date"] = date(2026, 1, 5)
at.session_state["cfg_calibration_offset_days"] = 2
at.session_state["cfg_w_capacity"] = 0.0
at.session_state["cfg_shutdown_weeks"] = "10,11"
at.run()
report(at)
"""))


def test_week_pickers_migrate_a_legacy_comma_string():
    """A session saved before the pickers existed must load, not crash.

    The shutdown-week setting used to be free text. The picker now holds a list,
    so anything stored as "10,11" has to be converted on sight — against the real
    multiselect, which rejects a string where it expects a sequence of options.
    """
    _assert_rendered(_run(_HEADER + """
from datetime import date

at = AppTest.from_file("app.py", default_timeout=300)
at.session_state["cfg_use_reference"] = True
at.session_state["cfg_reference_week_date"] = date(2026, 1, 5)
at.session_state["cfg_shutdown_weeks"] = "10,11"
at.session_state["cfg_curium_unavailable_weeks"] = "3"
at.run()
report(at)

migrated = at.session_state["cfg_shutdown_weeks"]
assert migrated == [10, 11], migrated
assert at.session_state["cfg_curium_unavailable_weeks"] == [3]
print("migrated:", migrated)
"""))


def test_onboarding_window_renders_as_dates_with_a_reference_week():
    """The window switches to date columns, and back, without losing rows."""
    _assert_rendered(_run(_HEADER + """
from datetime import date

import pandas as pd

at = AppTest.from_file("app.py", default_timeout=300)
at.session_state["cfg_use_reference"] = True
at.session_state["cfg_reference_week_date"] = date(2026, 1, 5)
at.session_state["ob_customers"] = pd.DataFrame([{
    "Site_ID": "NEW-1", "Site_Name": "Trial", "Earliest_Week": 3,
    "Latest_Week": 9, "Interval_Weeks": 7, "Country": "usa",
    "EU_Restricted": False,
}])
at.run()
report(at)

frame = at.session_state["ob_customers"]
assert "Earliest_Date" in frame.columns, list(frame.columns)
assert list(frame["Site_ID"]) == ["NEW-1"]
print("columns:", list(frame.columns))
"""))


def test_onboarding_window_stays_numeric_without_a_reference_week():
    _assert_rendered(_run(_HEADER + """
import pandas as pd

at = AppTest.from_file("app.py", default_timeout=300)
at.session_state["ob_customers"] = pd.DataFrame([{
    "Site_ID": "NEW-1", "Site_Name": "Trial", "Earliest_Week": 3,
    "Latest_Week": 9, "Interval_Weeks": 7, "Country": "usa",
    "EU_Restricted": False,
}])
at.run()
report(at)

frame = at.session_state["ob_customers"]
assert "Earliest_Week" in frame.columns, list(frame.columns)
assert "Earliest_Date" not in frame.columns
print("columns:", list(frame.columns))
"""))


def test_optimizer_result_panel_renders():
    """The result tables (plan, quota, changed weeks) use the real dataframe API."""
    _assert_rendered(_run(_HEADER + """
from datetime import date
import pandas as pd

from domain.delivery_assignment import DeliveryRecord
from domain.quota import QuarterlyQuotaStatus
from services.dtos import OptimizationResult
from ui import state as S

plan = pd.DataFrame({
    "Week": [1, 2],
    "Week_Type": ["Normal", "Partial"],
    "Demand_Due": [10, 12],
    "Good_Production": [15, 15],
    "Composite_Cost_USD": [1000.0, 2000.0],
})
result = OptimizationResult(
    plan_df=plan,
    summary={
        "total_composite_cost": 3000.0, "total_penalty_cost": 1000.0,
        "total_overtime_cost": 2000.0, "total_capacity_cost": 0.0,
        "overtime_weeks": 1, "total_quota_penalty_cost": 0.0,
        "partial_quarter_note": "Q1 3/13 wks fall partly outside the plan.",
    },
    issues_df=pd.DataFrame(columns=["row_index", "site_id", "issue"]),
    active_df=pd.DataFrame({"site_id": ["00449"]}),
    quota_status=[
        QuarterlyQuotaStatus(
            supplier="Curium", quarter=1, weeks=(1, 2), quota_mci=10000.0,
            ordered_mci=1500.0, remaining_mci=8500.0, shortfall_mci=0.0,
            penalty_usd=0.0, is_partial=True, weeks_covered=2,
            expected_weeks=13, prorated_quota_mci=1538.0,
            prorated_shortfall_mci=38.0, status="Partial — not penalised",
        ),
    ],
    assignments=[
        DeliveryRecord(
            site_id="00449", site_name="Alpha", country="usa", due_week=10,
            planned_week=8, due_week_shift=-2, manual_week=10, week_shift=-2,
            is_early=True, compared=True,
        ),
    ],
    change_summary={
        "compared": True, "total": 1, "unchanged": 0, "early": 1, "late": 0,
        "new_customers": 0, "uncomparable": 0,
    },
    week_dates=[(1, date(2026, 1, 5), date(2026, 1, 9)),
                (2, date(2026, 1, 12), date(2026, 1, 16))],
    xlsx_bytes=b"XLSX",
    warnings=["a warning"],
)

at = AppTest.from_file("app.py", default_timeout=300)
at.session_state[S.OPT_RESULT] = result
at.run()
report(at)
"""))


def test_converter_result_panel_renders():
    """The Import Manual Plan result tables render against the real API."""
    _assert_rendered(_run(_HEADER + """
import pandas as pd

from domain.site_derivation import (
    PlannerColumn, derive_sites, mapping_frame, notes_frame, sites_frame,
)
from services.dtos import ConversionResult
from ui import state as S

columns = [
    PlannerColumn("00449  Alpha Care, CA (7)", marks=(1, 8, 15), column_letter="L"),
    PlannerColumn("Beta Cardiology, TN", marks=(), column_letter="M"),
]
sites, warnings = derive_sites(columns, lambda h: "RF-0001")

at = AppTest.from_file("app.py", default_timeout=300)
at.session_state[S.CONV_RESULT] = ConversionResult(
    sites_df=sites_frame(sites),
    mapping_df=mapping_frame(sites),
    notes_df=notes_frame(sites),
    xlsx_bytes=b"XLSX",
    year=2026, site_count=2, active_count=1,
    generated_code_count=1, eu_restricted_count=0,
    scheduled_deliveries=3, implied_deliveries=8,
    issues_df=pd.DataFrame(columns=["row_index", "site_id", "issue"]),
    warnings=warnings + ["check the intervals"],
)
at.run()
report(at)
"""))


def test_comparison_result_panel_renders():
    """The Comparison tables render against the real API."""
    _assert_rendered(_run(_HEADER + """
import pandas as pd

from domain.comparison import BaselineResult
from io_adapters.master_planner_parser import AssignedId
from services.dtos import ComparisonResult
from ui import state as S

components = [
    {"Component": "Penalty", "Baseline": 2000.0, "Optimized": 500.0,
     "Saving_Abs": 1500.0, "Saving_Pct": 75.0},
    {"Component": "Total Composite", "Baseline": 3000.0, "Optimized": 1000.0,
     "Saving_Abs": 2000.0, "Saving_Pct": 66.7},
]

at = AppTest.from_file("app.py", default_timeout=300)
at.session_state[S.CMP_RESULT] = ComparisonResult(
    components=components,
    overtime_baseline=3, overtime_optimized=1,
    weekly_comparison=pd.DataFrame({
        "Week": [1, 2], "Manual_Production": [30, 20],
        "Optimized_Production": [25, 25],
    }),
    assigned_ids=[AssignedId("RF-0001", "Beta Cardiology, TN",
                             "Beta Cardiology, TN", "beta cardiology tn")],
    warnings=["a warning"],
)
at.run()
report(at)
"""))


# ---------------------------------------------------------------------------
# Files are loaded once, in the sidebar, and shared by every tab
# ---------------------------------------------------------------------------

def test_only_the_sidebar_uploads_files():
    """Four uploaders across four tabs became two, both in the sidebar.

    Each tab used to carry its own, and three separate Master Planner sheet
    dropdowns all wrote to one shared setting. Because Streamlit renders every tab
    on each interaction, the last dropdown to render won — so changing the sheet in
    Comparison silently left Onboarding parsing a summary sheet.
    """
    import pathlib
    import re

    uploaders = {}
    for path in pathlib.Path("ui").glob("*.py"):
        count = len(re.findall(r"st\.file_uploader\(", path.read_text()))
        if count:
            uploaders[path.name] = count

    assert uploaders == {"sidebar.py": 2}, uploaders


def test_only_one_widget_sets_each_sheet():
    """One owner per setting, so no tab can overwrite another's choice."""
    import pathlib
    import re

    writers = {}
    for path in pathlib.Path("ui").glob("*.py"):
        text = path.read_text()
        for setting in ("S.MP_SHEET", "S.SITES_SHEET"):
            hits = len(re.findall(rf"{re.escape(setting)}\]\s*=", text))
            if hits:
                writers.setdefault(setting, {})[path.name] = hits

    # tab_converter also seeds the sites sheet when handing its output over.
    assert set(writers["S.MP_SHEET"]) == {"sidebar.py"}, writers["S.MP_SHEET"]
    assert set(writers["S.SITES_SHEET"]) <= {"sidebar.py", "tab_converter.py"}, \
        writers["S.SITES_SHEET"]


def test_a_file_loaded_once_reaches_every_tab():
    """The sidebar writes shared state; no tab should ask for the file again."""
    _assert_rendered(_run(_HEADER + """
import io

import openpyxl
from streamlit.testing.v1 import AppTest

# Minimal sites workbook and Master Planner, built in-process.
sites = openpyxl.Workbook()
ws = sites.active
ws.title = "Sites"
ws.append(["Site_ID", "Active", "Next_Demand_Week", "Interval_Weeks", "Country"])
ws.append(["S1", "Y", 2, 7, "usa"])
sb = io.BytesIO(); sites.save(sb)

from datetime import datetime
mp = openpyxl.Workbook()
mw = mp.active
mw.title = "Schedule"
mw.append([None] * 12)
mw.append(["Weeks #", "MFG Date", "Calibration date", "Problem?", "Month", "FY",
           "US Demand", "RoW Demand", "Total Commercial", "QC GEN", "US STAB",
           "00449  Alpha Care, CA (7)"])
mw.append([1, datetime(2026, 1, 5), datetime(2026, 1, 9), None, "Jan", 2026,
           1, 0, 1, 1, None, 1])
mb = io.BytesIO(); mp.save(mb)

at = AppTest.from_file("app.py", default_timeout=300)
at.session_state["sites_bytes"] = sb.getvalue()
at.session_state["sites_name"] = "sites.xlsx"
at.session_state["sites_sheet_names"] = ["Sites"]
at.session_state["mp_bytes"] = mb.getvalue()
at.session_state["mp_name"] = "planner.xlsx"
at.session_state["mp_sheet_names"] = ["Schedule"]
at.run()
report(at)

# The sidebar resolved both sheets, and every tab rendered against them.
assert at.session_state["sites_sheet"] == "Sites", at.session_state["sites_sheet"]
assert at.session_state["mp_sheet"] == "Schedule", at.session_state["mp_sheet"]
print("sheets:", at.session_state["sites_sheet"], at.session_state["mp_sheet"])
"""))


# ---------------------------------------------------------------------------
# A manual plan on its own is enough to optimize
# ---------------------------------------------------------------------------

_PLANNER_FIXTURE = """
import io
from datetime import datetime

import openpyxl

def planner_bytes(weeks=12):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Schedule"
    ws.append([None] * 13)
    ws.append(["Weeks #", "MFG Date", "Calibration date", "Problem?", "Month",
               "FY", "US Demand", "RoW Demand", "Total Commercial", "QC GEN",
               "US STAB", "00449  Alpha Care, CA (7)", "Beta Imaging, TX (7)"])
    monday = datetime(2026, 1, 5)
    for week in range(1, weeks + 1):
        mfg = datetime.fromordinal(monday.toordinal() + 7 * (week - 1))
        a = 1 if (week - 1) % 7 == 0 else None
        b = 1 if (week - 3) % 7 == 0 and week >= 3 else None
        ws.append([week, mfg, datetime.fromordinal(mfg.toordinal() + 4), None,
                   mfg.strftime("%b"), 2026, 0, 0,
                   int(bool(a)) + int(bool(b)), 1, None, a, b])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
"""


def test_sites_can_be_derived_from_the_manual_plan_alone():
    """Uploading only the Master Planner must be enough to reach the optimizer."""
    _assert_rendered(_run(_HEADER + _PLANNER_FIXTURE + """
at = AppTest.from_file("app.py", default_timeout=300)
at.session_state["mp_bytes"] = planner_bytes()
at.session_state["mp_name"] = "planner.xlsx"
at.session_state["mp_sheet_names"] = ["Schedule"]
at.run()
report(at)

# No sites file yet, so the tabs must point at the derive action rather than
# simply refusing.
assert at.session_state["sites_bytes"] is None
labels = [b.label for b in at.button]
assert any("Derive sites from the manual plan" in l for l in labels), labels
print("derive offered:", [l for l in labels if "Derive" in l])
"""))


def test_deriving_loads_a_sites_file_and_offers_it_for_download():
    """The planner asked to be able to keep the generated input file."""
    _assert_rendered(_run(_HEADER + _PLANNER_FIXTURE + """
at = AppTest.from_file("app.py", default_timeout=300)
at.session_state["mp_bytes"] = planner_bytes()
at.session_state["mp_name"] = "planner.xlsx"
at.session_state["mp_sheet_names"] = ["Schedule"]
at.run()
report(at)

derive = next(b for b in at.button if "Derive sites" in b.label)
derive.click().run()
report(at)

assert at.session_state["sites_bytes"], "deriving should load a sites file"
assert at.session_state["sites_from_planner"] is True
assert at.session_state["sites_name"] == "sites_from_manual_plan.xlsx"

# And it must be downloadable, not only usable in place.
downloads = [d.label for d in at.download_button]
assert any("Download this input file" in l for l in downloads), downloads
print("derived", len(at.session_state["sites_bytes"]), "bytes; downloads:", downloads)
"""))


def test_a_derived_sites_file_reaches_the_optimizer():
    _assert_rendered(_run(_HEADER + _PLANNER_FIXTURE + """
at = AppTest.from_file("app.py", default_timeout=300)
at.session_state["mp_bytes"] = planner_bytes()
at.session_state["mp_name"] = "planner.xlsx"
at.session_state["mp_sheet_names"] = ["Schedule"]
at.run()
next(b for b in at.button if "Derive sites" in b.label).click().run()
report(at)

# The sheet the optimizer will read is the derived workbook's Sites sheet.
assert at.session_state["sites_sheet"] == "Sites", at.session_state["sites_sheet"]
# The dropdown re-registers after the rerun, and must agree rather than
# reasserting a sheet name belonging to the file it replaced.
assert at.session_state["sb_sites_sheet"] == "Sites", \\
    at.session_state["sb_sites_sheet"]
print("optimizer would read sheet:", at.session_state["sites_sheet"])
"""))
