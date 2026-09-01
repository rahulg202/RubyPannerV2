"""Tests for writing site codes into a Master Planner copy.

The critical property is non-destructiveness: the planner's workbook is their
source of truth, so the annotator must return a copy, must not touch columns that
already carry an account number, and must not strip formulas out of the sheet.
"""

from __future__ import annotations

import io

import openpyxl
import pytest

from domain.site_derivation import PlannerColumn, derive_sites
from io_adapters.master_planner_annotator import (
    MasterPlannerAnnotator,
    annotate_master_planner,
)
from io_adapters.master_planner_parser import assign_stable_id

AGGREGATES = [
    "Weeks #", "MFG Date", "Calibration date", "Problem?", "Month", "FY",
    "US Demand", "RoW Demand", "Total Commercial", "QC GEN", "US STAB",
]
# Column L is numbered, M and N are not.
CUSTOMERS = [
    "00449    Alpha Specialty Care, Fresno, CA (7)",
    "Beta Cardiology, Jackson, TN (7)",
    "Gamma Imaging, Portland, OR (7)",
]


def _workbook(with_formula: bool = True) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Schedule"
    ws.append([None] * (len(AGGREGATES) + len(CUSTOMERS)))
    ws.append(AGGREGATES + CUSTOMERS)
    ws.append([1, None, None, None, "Jan", 2026, 0, 0,
               "=SUM(L3:N3)" if with_formula else 1, 1, None, 1, 1, None])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _sites():
    columns = [
        PlannerColumn(CUSTOMERS[0], marks=(1,), column_letter="L"),
        PlannerColumn(CUSTOMERS[1], marks=(1,), column_letter="M"),
        PlannerColumn(CUSTOMERS[2], marks=(), column_letter="N"),
    ]
    # Use the real generator: the codes it produces must be the codes the parser
    # reads back, so a stand-in of a different shape would not exercise that.
    sites, _ = derive_sites(columns, assign_stable_id)
    return sites


def _headers(data: bytes) -> dict[str, str]:
    wb = openpyxl.load_workbook(io.BytesIO(data))
    try:
        ws = wb["Schedule"]
        return {c.column_letter: c.value for c in ws[2] if c.value}
    finally:
        wb.close()


@pytest.fixture
def annotated():
    return MasterPlannerAnnotator().annotate(_workbook(), _sites())


def test_generated_codes_are_written_into_the_headers(annotated):
    headers = _headers(annotated.xlsx_bytes)
    assert headers["M"].startswith("RF-")
    assert headers["N"].startswith("RF-")
    assert "Beta Cardiology" in headers["M"]


def test_numbered_columns_are_left_alone(annotated):
    headers = _headers(annotated.xlsx_bytes)
    assert headers["L"] == CUSTOMERS[0]


def test_only_generated_codes_are_counted(annotated):
    assert annotated.written_count == 2
    assert {col for _code, col in annotated.written} == {"M", "N"}


def test_the_original_workbook_is_not_modified():
    original = _workbook()
    MasterPlannerAnnotator().annotate(original, _sites())
    assert _headers(original)["M"] == CUSTOMERS[1]


def test_formulas_survive_the_round_trip(annotated):
    """Opening with cached values would silently strip every formula."""
    wb = openpyxl.load_workbook(io.BytesIO(annotated.xlsx_bytes))
    try:
        assert wb["Schedule"]["I3"].value == "=SUM(L3:N3)"
    finally:
        wb.close()


def test_running_twice_does_not_double_prefix():
    sites = _sites()
    first = MasterPlannerAnnotator().annotate(_workbook(), sites)
    second = MasterPlannerAnnotator().annotate(first.xlsx_bytes, sites)

    headers = _headers(second.xlsx_bytes)
    assert headers["M"].count("RF-") == 1
    assert second.written_count == 0
    assert len(second.already_present) == 2


def test_annotated_workbook_still_parses(annotated):
    from io_adapters.master_planner_parser import parse_master_planner

    parsed = parse_master_planner(annotated.xlsx_bytes, horizon_weeks=5, year=None)
    # The previously-unnumbered column now matches on its literal code, not a hash.
    assert any(k.startswith("RF-") for k in parsed.customer_schedule)


def test_annotation_preserves_the_codes_it_wrote(annotated):
    """Regression: annotating used to change the very codes it was writing.

    A header with 'RF-...' prepended has no leading digits, so the parser fell
    through to hashing the header text — which had just changed. The result was
    a workbook whose codes matched nothing.
    """
    from io_adapters.master_planner_parser import parse_master_planner

    before = parse_master_planner(_workbook(), horizon_weeks=5, year=None)
    after = parse_master_planner(annotated.xlsx_bytes, horizon_weeks=5, year=None)

    generated = {k for k in before.customer_schedule if k.startswith("RF-")}
    assert generated, "fixture should produce hash-derived codes to begin with"
    assert generated <= set(after.customer_schedule)


def test_codes_survive_an_edit_to_the_surrounding_header_text(annotated):
    """The reason for annotating at all: a typo fix must not change the code."""
    from io_adapters.master_planner_parser import parse_master_planner

    wb = openpyxl.load_workbook(io.BytesIO(annotated.xlsx_bytes))
    ws = wb["Schedule"]
    code = None
    for cell in ws[2]:
        if cell.value and str(cell.value).startswith("RF-"):
            code = str(cell.value).split()[0]
            cell.value = str(cell.value) + " (renamed site)"
            break
    buffer = io.BytesIO()
    wb.save(buffer)
    wb.close()

    parsed = parse_master_planner(buffer.getvalue(), horizon_weeks=5, year=None)
    assert code in parsed.customer_schedule


def test_site_without_a_column_reference_is_skipped_not_crashed():
    columns = [PlannerColumn("Delta Clinic, TX (7)", marks=(1,), column_letter="")]
    sites, _ = derive_sites(columns, lambda h: "RF-9999")
    result = MasterPlannerAnnotator().annotate(_workbook(), sites)
    assert result.written_count == 0
    assert result.skipped and result.skipped[0][0] == "RF-9999"


def test_missing_sheet_is_rejected():
    with pytest.raises(ValueError, match="not found"):
        MasterPlannerAnnotator().annotate(_workbook(), _sites(), sheet="Nope")


def test_module_level_wrapper():
    result = annotate_master_planner(_workbook(), _sites())
    assert result.xlsx_bytes and result.written_count == 2
