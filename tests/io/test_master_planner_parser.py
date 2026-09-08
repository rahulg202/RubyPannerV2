"""Tests for the Master Planner parser (io_adapters.master_planner_parser).

Uses a synthetic fixture workbook mirroring the real sheet's layout so the tests
do not depend on the (large, changing) production workbook.
"""

import io

import openpyxl
import pytest

from io_adapters.master_planner_parser import (
    MasterPlannerParser,
    assign_stable_id,
    parse_master_planner,
)


def _fixture_workbook() -> bytes:
    """Build a small Master Planner-shaped workbook.

    Row 1 blank-ish, row 2 headers, rows 3+ data. Customer cells hold 1 to mark
    a scheduled generator.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Schedule"
    ws.append([None] * 12)  # row 1 (spacer, like the real sheet)
    ws.append([
        "Weeks #", "MFG Date \n(Holidays)", "Calibration date", "Problem?",
        "Month", "FY", "US Demand", "RoW Demand", "Total Commercial", "QC GEN",
        "US\nSTAB",                                  # non-customer marker
        "00449    Acme Cardiology, Fresno, CA",      # numbered customer
        "Apex Cardiology, Jackson, TN (7 weeks)",    # unnumbered customer
    ])
    from datetime import datetime
    rows = [
        (1, datetime(2026, 1, 5), datetime(2026, 1, 9), None, "Jan", 2026, 2, 0, 2, 1, None, 1, 1),
        (2, datetime(2026, 1, 12), datetime(2026, 1, 16), None, "Jan", 2026, 1, 0, 1, 0, None, 1, None),
        (3, datetime(2026, 1, 19), datetime(2026, 1, 23), None, "Jan", 2026, 1, 0, 1, 1, None, None, 1),
    ]
    for r in rows:
        ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def fixture_bytes():
    return _fixture_workbook()


# ---------------------------------------------------------------------------
# Stable ID assignment
# ---------------------------------------------------------------------------

def test_stable_id_is_deterministic():
    h = "Apex Cardiology, Jackson, TN (7 weeks)"
    assert assign_stable_id(h) == assign_stable_id(h)


def test_stable_id_format():
    sid = assign_stable_id("Some Customer, TX (7)")
    assert sid.startswith("RF-") and len(sid) == 11


def test_stable_id_ignores_punctuation_and_case():
    a = assign_stable_id("Apex Cardiology, Jackson, TN (7)")
    b = assign_stable_id("apex cardiology  jackson tn 7")
    assert a == b


def test_stable_id_differs_for_different_customers():
    assert assign_stable_id("Customer A") != assign_stable_id("Customer B")


# ---------------------------------------------------------------------------
# Parsing
# ---------------------------------------------------------------------------

def test_parses_weeks_and_aggregates(fixture_bytes):
    d = parse_master_planner(fixture_bytes, horizon_weeks=52)
    assert d.weekly_commercial[1] == 2
    assert d.weekly_commercial[2] == 1
    assert d.weekly_qc[1] == 1
    assert d.weekly_qc[3] == 1


def test_planned_production_adds_qc(fixture_bytes):
    d = parse_master_planner(fixture_bytes, horizon_weeks=52)
    # week 1: commercial 2 + QC 1 = 3
    assert d.weekly_planned_production[1] == 3
    assert d.weekly_planned_production[2] == 1   # QC 0


def test_dates_extracted(fixture_bytes):
    from datetime import date
    d = parse_master_planner(fixture_bytes, horizon_weeks=52)
    assert d.mfg_dates[1] == date(2026, 1, 5)
    assert d.cal_dates[1] == date(2026, 1, 9)


def test_numbered_customer_matched_by_leading_number(fixture_bytes):
    d = parse_master_planner(fixture_bytes, horizon_weeks=52)
    assert "00449" in d.customer_schedule
    assert d.customer_schedule["00449"][1] == 1
    assert d.customer_schedule["00449"][2] == 1
    assert d.customer_schedule["00449"][3] == 0


def test_unnumbered_customer_gets_generated_id(fixture_bytes):
    d = parse_master_planner(fixture_bytes, horizon_weeks=52)
    gen = [a for a in d.assigned_ids
           if a.customer_name.startswith("Apex Cardiology")]
    assert len(gen) == 1
    sid = gen[0].generated_id
    assert sid in d.customer_schedule
    assert d.customer_schedule[sid][1] == 1
    assert d.customer_schedule[sid][3] == 1


def test_non_customer_columns_ignored(fixture_bytes):
    d = parse_master_planner(fixture_bytes, horizon_weeks=52)
    assert any("STAB" in c for c in d.ignored_columns)
    # "Problem?" is an aggregate header, not a customer
    assert not any(a.customer_name == "Problem?" for a in d.assigned_ids)


def test_customer_marks_reconcile_with_total_commercial(fixture_bytes):
    """The count of schedule marks must equal Total Commercial per week."""
    d = parse_master_planner(fixture_bytes, horizon_weeks=52)
    for wk in (1, 2, 3):
        marks = sum(sched[wk] for sched in d.customer_schedule.values())
        assert marks == d.weekly_commercial[wk]


def test_weeks_outside_horizon_excluded(fixture_bytes):
    d = parse_master_planner(fixture_bytes, horizon_weeks=2)
    assert 3 not in d.week_to_row
    assert d.rows_excluded >= 1


def test_a_wrong_sheet_falls_back_to_the_schedule_grid(fixture_bytes):
    """A planner pointed at the wrong sheet should still get a result.

    The real workbook has nine sheets and only one is the schedule grid. Failing
    on a sheet that could never have worked produced an error the planner could
    not act on, so the reader now finds the sheet that qualifies and says so.
    """
    d = parse_master_planner(fixture_bytes, sheet="NoSuchSheet")
    assert d.customer_schedule                      # it read the real grid
    assert any("Read 'Schedule' instead" in i for i in d.issues)


def test_a_workbook_with_no_schedule_grid_is_rejected_with_the_sheet_names():
    """When no sheet qualifies, say so and list what was actually there."""
    import io as _io

    import openpyxl as _openpyxl

    wb = _openpyxl.Workbook()
    wb.active.title = "Summary"
    wb.active.append(["Not", "a", "schedule"])
    wb.create_sheet("Notes")
    buffer = _io.BytesIO()
    wb.save(buffer)

    with pytest.raises(ValueError) as exc:
        parse_master_planner(buffer.getvalue())
    message = str(exc.value)
    assert "Weeks #" in message
    assert "Summary" in message and "Notes" in message


def test_week_to_row_points_at_real_spreadsheet_rows(fixture_bytes):
    """Empty rows are skipped while streaming, so this must not count positions."""
    d = parse_master_planner(fixture_bytes, horizon_weeks=52)
    # Fixture data starts on row 3, one week per row.
    assert d.week_to_row[1] == 3
    assert d.week_to_row[2] == 4


def test_parser_class_and_function_agree(fixture_bytes):
    a = MasterPlannerParser().parse(fixture_bytes, "Schedule", 52, None)
    b = parse_master_planner(fixture_bytes, "Schedule", 52, None)
    assert a.weekly_planned_production == b.weekly_planned_production
    assert a.customer_schedule == b.customer_schedule


# ---------------------------------------------------------------------------
# Codes already written into a header (annotated workbooks)
# ---------------------------------------------------------------------------

def test_explicit_generated_code_in_a_header_is_used_verbatim():
    """An annotated workbook must yield the code in the header, not a fresh hash.

    Regression: the parser looked only for leading digits, so a header starting
    'RF-...' fell through to hashing the header text. Annotating a workbook then
    produced codes that differed from the ones written into it.
    """
    import io as _io
    from datetime import datetime

    import openpyxl as _openpyxl

    header = "RF-4e40dfd2  Apex Cardiology, Jackson, TN (7 weeks)"
    wb = _openpyxl.Workbook()
    ws = wb.active
    ws.title = "Schedule"
    ws.append([None] * 12)
    ws.append([
        "Weeks #", "MFG Date", "Calibration date", "Problem?", "Month", "FY",
        "US Demand", "RoW Demand", "Total Commercial", "QC GEN", "US STAB",
        header,
    ])
    ws.append([1, datetime(2026, 1, 5), datetime(2026, 1, 9), None, "Jan", 2026,
               1, 0, 1, 1, None, 1])
    buffer = _io.BytesIO()
    wb.save(buffer)

    parsed = parse_master_planner(buffer.getvalue(), horizon_weeks=5)

    assert "RF-4e40dfd2" in parsed.customer_schedule
    assert parsed.customer_schedule["RF-4e40dfd2"][1] == 1
    # It is an explicit identifier, so it is not reported as newly assigned.
    assert not any(a.generated_id == "RF-4e40dfd2" for a in parsed.assigned_ids)
    assert parsed.customer_headers["RF-4e40dfd2"] == header


# ---------------------------------------------------------------------------
# Streaming: empty padding rows must not be retained
# ---------------------------------------------------------------------------

def _padded_workbook(padding_rows: int = 3000) -> bytes:
    """A small schedule followed by thousands of empty rows.

    The real planner extends to 15,300 rows of which only ~500 carry a week
    number. Holding a cell object per cell cost 2 GB of peak memory and got the
    hosted app killed, so the reader must stream and discard.
    """
    import io as _io
    from datetime import datetime

    import openpyxl as _openpyxl

    wb = _openpyxl.Workbook()
    ws = wb.active
    ws.title = "Schedule"
    ws.append([None] * 12)
    ws.append([
        "Weeks #", "MFG Date", "Calibration date", "Problem?", "Month", "FY",
        "US Demand", "RoW Demand", "Total Commercial", "QC GEN", "US STAB",
        "00449    Alpha Care, CA (7)",
    ])
    monday = datetime(2026, 1, 5)
    for week in range(1, 4):
        mfg = datetime.fromordinal(monday.toordinal() + 7 * (week - 1))
        ws.append([week, mfg, mfg, None, "Jan", 2026, 1, 0, 1, 1, None, 1])
    for _ in range(padding_rows):
        ws.append([None] * 12)

    buffer = _io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_empty_padding_rows_are_discarded_while_streaming():
    data = _padded_workbook(3000)
    parsed = parse_master_planner(data, horizon_weeks=52)

    # Only the three real weeks are kept, regardless of how far the sheet runs.
    assert sorted(parsed.week_to_row) == [1, 2, 3]
    assert parsed.rows_excluded == 0
    assert sum(parsed.customer_schedule["00449"]) == 3


def test_padding_does_not_shift_the_row_mapping():
    data = _padded_workbook(500)
    parsed = parse_master_planner(data, horizon_weeks=52)
    assert parsed.week_to_row[1] == 3      # real spreadsheet row, not a position
